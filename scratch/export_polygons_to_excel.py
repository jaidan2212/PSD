import os
import glob
import math
import shapefile
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_excel():
    excel_path = r"c:\laragon\www\TugasPSD\Data_Poligon_dan_Piksel_LULC_Jatim.xlsx"
    csv_poligon_path = r"c:\laragon\www\TugasPSD\Data_Poligon_Shapefile_Jatim.csv"
    
    # 1. Load CSV data
    df_pixels = pd.read_csv(r"c:\laragon\www\TugasPSD\dataset_sentinel2_jatim.csv")
    
    # 2. Extract Shapefiles
    shp_mapping = [
        {"kelas_id": 1, "kelas": "Sawah", "shp": "sawah", "file_zip": "50 Sawah gqis.qgz / sawah.zip", "lokasi": "Lumbung padi Jatim (Sidoarjo, Lamongan, Jombang, Nganjuk, Malang)"},
        {"kelas_id": 2, "kelas": "Bangunan", "shp": "bangunan", "file_zip": "50 Non Sawah gqis.qgz", "lokasi": "Permukiman & infrastruktur perkotaan (Kamal Bangkalan, Madura & Surabaya)"},
        {"kelas_id": 3, "kelas": "Mangrove", "shp": "mangrove", "file_zip": "Mangrove Zaidan.zip", "lokasi": "Pesisir hutan bakau Jatim (Wonorejo Surabaya, pesisir Madura, Banyuwangi)"},
        {"kelas_id": 4, "kelas": "Lahan Hijau", "shp": "lahan_hijau", "file_zip": "Lahan hijau.zip", "lokasi": "Hutan & vegetasi perbukitan utara Jatim (kawasan Tuban, Lamongan)"},
        {"kelas_id": 5, "kelas": "Perairan Terbuka (Laut)", "shp": "lautan", "file_zip": "laut.zip", "lokasi": "Perairan Laut Jawa (utara) & Samudra Hindia (selatan) sekeliling Jatim"},
        {"kelas_id": 6, "kelas": "Danau", "shp": "danau", "file_zip": "Danau.zip", "lokasi": "Perairan tawar alami/buatan Jatim (Waduk Karangkates, Selorejo, Ranu)"},
    ]
    
    list_poligon_detail = []
    no = 1
    
    for item in shp_mapping:
        shp_dir = os.path.join(r"c:\laragon\www\TugasPSD\shapefiles", item["shp"])
        shp_files = glob.glob(os.path.join(shp_dir, "*.shp"))
        if not shp_files:
            continue
        sf = shapefile.Reader(shp_files[0])
        
        for idx, sr in enumerate(sf.shapeRecords()):
            s = sr.shape
            bbox = s.bbox # xmin, ymin, xmax, ymax
            xmin, ymin, xmax, ymax = bbox[0], bbox[1], bbox[2], bbox[3]
            c_lon = (xmin + xmax) / 2.0
            c_lat = (ymin + ymax) / 2.0
            pts_cnt = len(s.points)
            
            list_poligon_detail.append({
                "No": no,
                "Kelas_ID": item["kelas_id"],
                "Kelas_Lahan": item["kelas"],
                "Poligon_ID_Shapefile": idx,
                "Centroid_Lon": round(c_lon, 6),
                "Centroid_Lat": round(c_lat, 6),
                "Min_Lon": round(xmin, 6),
                "Max_Lon": round(xmax, 6),
                "Min_Lat": round(ymin, 6),
                "Max_Lat": round(ymax, 6),
                "Jumlah_Titik_Verteks": pts_cnt,
                "File_Shapefile_Asal": f"{item['shp']}.shp",
                "Sumber_Data_QGIS": item["file_zip"]
            })
            no += 1
            
    df_polygons = pd.DataFrame(list_poligon_detail)
    df_polygons.to_csv(csv_poligon_path, index=False)
    
    # 3. Create Ringkasan Kelas LULC
    ringkasan_list = []
    for item in shp_mapping:
        poly_sub = df_polygons[df_polygons["Kelas_ID"] == item["kelas_id"]]
        pix_sub = df_pixels[df_pixels["kelas_id"] == item["kelas_id"]]
        
        ringkasan_list.append({
            "ID Kelas": item["kelas_id"],
            "Nama Kelas Tutupan Lahan": item["kelas"],
            "Jumlah Poligon ROI (QGIS)": len(poly_sub),
            "Jumlah Sampel Piksel Diekstrak": len(pix_sub),
            "Resolusi Spasial Sensor": "10 m x 10 m (1 piksel = 100 m²)",
            "Rentang Bujur (Longitude)": f"{pix_sub['lon'].min():.4f} s/d {pix_sub['lon'].max():.4f}",
            "Rentang Lintang (Latitude)": f"{pix_sub['lat'].min():.4f} s/d {pix_sub['lat'].max():.4f}",
            "File Sumber Shapefile": item["file_zip"],
            "Area Wilayah Geografis di Jawa Timur": item["lokasi"]
        })
    df_ringkasan = pd.DataFrame(ringkasan_list)
    
    # 4. Agregasi Piksel per Poligon
    df_agg = df_pixels.groupby(["kelas_id", "kelas", "poligon_id"]).agg(
        Jumlah_Piksel=("lon", "count"),
        Rata_Lon=("lon", "mean"),
        Rata_Lat=("lat", "mean"),
        Rata_NDVI=("NDVI", "mean"),
        Rata_NDWI=("NDWI", "mean"),
        Rata_MNDWI=("MNDWI", "mean"),
        Rata_NDBI=("NDBI", "mean"),
        Rata_EVI=("EVI", "mean")
    ).reset_index().round(4)
    
    # 5. Write to Multi-Sheet Excel with openpyxl
    wb = openpyxl.Workbook()
    # remove default sheet
    wb.remove(wb.active)
    
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid") # Dark Blue
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    border_thin = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    def write_sheet(title, df_data, col_widths=None):
        ws = wb.create_sheet(title=title)
        ws.views.sheetView[0].showGridLines = True
        
        # Header
        cols = list(df_data.columns)
        ws.append(cols)
        for col_idx in range(1, len(cols) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border_thin
        ws.row_dimensions[1].height = 28
        
        # Rows
        for r_idx, row in enumerate(df_data.itertuples(index=False), start=2):
            ws.append(list(row))
            for c_idx in range(1, len(cols) + 1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.font = data_font
                cell.border = border_thin
                val = cell.value
                if isinstance(val, (int, float)):
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
            ws.row_dimensions[r_idx].height = 20
            
        # Adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
            
        return ws
    
    write_sheet("Ringkasan_Kelas_LULC", df_ringkasan)
    write_sheet("Detail_210_Poligon_Shapefile", df_polygons)
    write_sheet("Piksel_per_Poligon_ROI", df_agg)
    write_sheet("Data_Piksel_Sentinel2_520", df_pixels.round(4))
    
    wb.save(excel_path)
    print(f"[SUCCESS] File Excel berhasil dibuat di: {excel_path}")
    print(f"[SUCCESS] File CSV poligon berhasil dibuat di: {csv_poligon_path}")

if __name__ == "__main__":
    generate_excel()
